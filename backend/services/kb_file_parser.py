import os
import re
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None


MAX_KB_FILE_TEXT_CHARS = int(os.getenv("KB_FILE_TEXT_MAX_CHARS", "40000"))

# 旧版 Office 二进制格式 → 新版 OOXML 格式
LEGACY_OFFICE_TARGETS = {".doc": "docx", ".ppt": "pptx", ".xls": "xlsx"}

# 图片 → 多模态模型描述
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".gif"}
IMAGE_DESCRIBE_PROMPT = (
    "请用中文详细描述这张图片的内容，包括文字、表格、图标、场景、人物、物体、颜色、布局等所有可识别的信息。"
    "如果图中有表格或文字，请逐项列出。描述应尽量完整，后续会基于这段描述做检索。"
)


def describe_image_via_vllm(path: Path) -> str:
    """用 vLLM 的多模态接口把图片转成中文描述，作为可被向量检索的文本块。"""
    if not path.exists():
        return ""
    try:
        import base64, mimetypes
        from openai import OpenAI
    except Exception:
        return "图片解析失败：依赖不可用"
    mime, _ = mimetypes.guess_type(path.name)
    if not mime or not mime.startswith("image/"):
        mime = "image/jpeg"
    try:
        b64 = base64.b64encode(path.read_bytes()).decode("ascii")
    except Exception as exc:
        return f"图片读取失败：{exc}"
    data_url = f"data:{mime};base64,{b64}"
    base_url = os.getenv("OPENAI_API_BASE") or os.getenv("OPENAI_BASE_URL") or "http://127.0.0.1:62272/v1"
    api_key = os.getenv("OPENAI_API_KEY", "EMPTY")
    model_name = os.getenv("CHAT_MODEL_NAME", "Qwen3.5-27B")
    try:
        client = OpenAI(base_url=base_url, api_key=api_key, timeout=float(os.getenv("CHAT_MODEL_TIMEOUT", "120")))
        resp = client.chat.completions.create(
            model=model_name,
            temperature=0,
            max_tokens=int(os.getenv("CHAT_MODEL_MAX_TOKENS", "4096")),
            extra_body={"chat_template_kwargs": {"enable_thinking": False}},
            messages=[{
                "role": "user",
                "content": [
                    {"type": "text", "text": IMAGE_DESCRIBE_PROMPT},
                    {"type": "image_url", "image_url": {"url": data_url}},
                ],
            }],
        )
        text = (resp.choices[0].message.content or "").strip()
        header = f"【图片文件：{path.name}】\n"
        return compact_text(header + text)
    except Exception as exc:
        return f"图片描述生成失败：{exc}"


def _find_soffice() -> str:
    """优先找 soffice；找不到再试 libreoffice。"""
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    return ""


def convert_legacy_office(source: Path, target_ext: str) -> Path | None:
    """用 LibreOffice headless 把旧版 Office 转成 OOXML。
    返回新文件的 Path（位于临时目录），调用方用完自己删；失败返回 None。
    并发安全：每次调用独立 UserInstallation profile。
    """
    soffice = _find_soffice()
    if not soffice or not source.exists():
        return None
    tmp_root = Path(tempfile.gettempdir())
    out_dir = tmp_root / f"lo_out_{uuid.uuid4().hex}"
    profile = tmp_root / f"lo_profile_{uuid.uuid4().hex}"
    out_dir.mkdir(parents=True, exist_ok=True)
    try:
        cmd = [
            soffice, "--headless",
            f"-env:UserInstallation=file://{profile}",
            "--convert-to", target_ext,
            "--outdir", str(out_dir),
            str(source),
        ]
        proc = subprocess.run(cmd, capture_output=True, timeout=60)
        if proc.returncode != 0:
            return None
        # 输出文件名 = source.stem + target_ext
        expected = out_dir / f"{source.stem}.{target_ext}"
        if expected.exists():
            return expected
        # 兜底：扫 out_dir 取第一个 target_ext 文件
        for f in out_dir.glob(f"*.{target_ext}"):
            return f
        return None
    except Exception:
        return None
    finally:
        shutil.rmtree(profile, ignore_errors=True)


def _extract_with_legacy_conversion(source: Path, parser) -> str:
    """把旧版 Office 转成新版再调传入的 parser。"""
    target_ext = LEGACY_OFFICE_TARGETS.get(source.suffix.lower())
    if not target_ext:
        return ""
    converted = convert_legacy_office(source, target_ext)
    if not converted:
        return "该文件为旧版 Office 格式，服务器未能成功转换，请手动转换为 docx/xlsx/pptx 后重试。"
    try:
        return parser(converted)
    finally:
        shutil.rmtree(converted.parent, ignore_errors=True)


def compact_text(text: str, limit: int = MAX_KB_FILE_TEXT_CHARS) -> str:
    # 保留换行结构（表格/PDF 分行），只压缩行内多余空白
    src = str(text or "")
    lines = [re.sub(r"[ \t\r\f\v]+", " ", line).strip() for line in src.splitlines()]
    compact = "\n".join(line for line in lines if line)
    if len(compact) <= limit:
        return compact
    return compact[:limit] + " ...(已截断)"


def _extract_text_nodes_from_xml(xml_bytes: bytes) -> list[str]:
    try:
        root = ET.fromstring(xml_bytes)
    except Exception:
        return []
    texts = []
    for elem in root.iter():
        if elem.tag.endswith('}t') or elem.tag == 't':
            if elem.text and elem.text.strip():
                texts.append(elem.text.strip())
    return texts


def extract_pdf_text(path: Path) -> str:
    parts = []
    if PdfReader is not None:
        try:
            reader = PdfReader(str(path))
            for page in reader.pages:
                parts.append(page.extract_text() or "")
        except Exception:
            parts = []
    if not ''.join(parts).strip():
        try:
            result = subprocess.run(["pdftotext", str(path), "-"], capture_output=True, text=True, check=True)
            return compact_text(result.stdout)
        except Exception:
            return ""
    return compact_text("\n".join(parts))


def extract_docx_text(path: Path) -> str:
    try:
        with zipfile.ZipFile(path) as zf:
            data = zf.read('word/document.xml')
    except Exception:
        return ''
    return compact_text("\n".join(_extract_text_nodes_from_xml(data)))


def extract_pptx_text(path: Path) -> str:
    texts = []
    try:
        with zipfile.ZipFile(path) as zf:
            for name in sorted(zf.namelist()):
                if name.startswith('ppt/slides/slide') and name.endswith('.xml'):
                    texts.extend(_extract_text_nodes_from_xml(zf.read(name)))
    except Exception:
        return ''
    return compact_text("\n".join(texts))


def extract_xlsx_text(path: Path) -> str:
    """按行读取 xlsx，每行一条，保留行列结构，方便模型理解表格。"""
    try:
        from openpyxl import load_workbook
    except Exception:
        return ''
    try:
        wb = load_workbook(str(path), data_only=True, read_only=True)
    except Exception:
        return ''
    parts: list[str] = []
    try:
        for sheet in wb.worksheets:
            parts.append(f'【Sheet: {sheet.title}】')
            for row in sheet.iter_rows(values_only=True):
                cells = ['' if v is None else str(v).strip() for v in row]
                if not any(cells):
                    continue
                parts.append(' | '.join(cells))
    finally:
        wb.close()
    return compact_text('\n'.join(parts))


def extract_kb_file_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    try:
        if suffix in {'.txt', '.md', '.csv', '.json', '.py', '.log'}:
            return compact_text(file_path.read_text(encoding='utf-8', errors='ignore'))
        if suffix == '.pdf':
            return extract_pdf_text(file_path)
        if suffix == '.docx':
            return extract_docx_text(file_path)
        if suffix == '.pptx':
            return extract_pptx_text(file_path)
        if suffix == '.xlsx':
            return extract_xlsx_text(file_path)
        if suffix == '.doc':
            return _extract_with_legacy_conversion(file_path, extract_docx_text)
        if suffix == '.xls':
            return _extract_with_legacy_conversion(file_path, extract_xlsx_text)
        if suffix == '.ppt':
            return _extract_with_legacy_conversion(file_path, extract_pptx_text)
        if suffix in IMAGE_SUFFIXES:
            return describe_image_via_vllm(file_path)
        return '该文件类型暂不支持正文抽取，但文件已上传保存。'
    except Exception as exc:
        return f'文件解析失败：{exc}'
