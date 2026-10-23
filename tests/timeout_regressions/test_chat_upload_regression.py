"""In-process tests of original /api/upload -> /api/chat route bodies.

XTY_SOURCE_ROOT selects a staged or production tree without importing app.py.
Real MD/DOCX/XLSX/PDF parsers run on synthetic fixtures. Storage, identity,
history persistence, and model/recommendation generation are memory substitutes.
"""

import ast
import asyncio
import base64
import contextlib
import hashlib
import io
import json
import logging
import mimetypes
import os
from pathlib import Path
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import types
import unittest
from unittest.mock import patch
import uuid
import xml.etree.ElementTree as ET
import zipfile
from typing import Any, Dict, List, Optional

from docx import Document
from fastapi import FastAPI, File, Request, UploadFile
from fastapi.responses import JSONResponse, StreamingResponse
import httpx
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


ROOT = Path(os.environ.get("XTY_SOURCE_ROOT", str(Path(__file__).resolve().parents[1])))
SOURCE = ROOT / "backend/app.py"
REPORT = Path(os.environ.get("XTY_CHAT_TEST_REPORT", str(Path(__file__).resolve().parents[1] / "diagnostics/chat_upload_results.json")))
FIXED_ANSWER = "Offline model fixture response."
FUNCTION_NAMES = {
    "success_response", "error_response", "normalize_db_version",
    "_compact_text", "_extract_pdf_text", "_extract_text_nodes_from_xml",
    "_extract_docx_text", "_extract_pptx_text", "_extract_xlsx_text",
    "extract_uploaded_file_text", "build_uploaded_file_contexts",
    "_strip_rag_sources", "compose_chat_prompt", "_sse_event", "_sse_comment",
    "_register_stream", "_unregister_stream", "_build_chat_done_payload",
    "_get_request_payload", "chat_endpoint", "upload_file",
    "resolve_staging_files", "move_staging_to_dest",
    "_decode_graph_local_file_id",
}
CASE_DETAILS = []


class MemoryStorage:
    def __init__(self):
        self.objects = {}
        self.fail_upload = False
        self.fail_download = False
        self.fail_move = False
        self.operations = {name: 0 for name in ("upload", "list", "move", "download")}

    def upload_file_obj(self, file_obj, object_name, content_type=None):
        self.operations["upload"] += 1
        if self.fail_upload:
            return False
        self.objects[object_name] = file_obj.read()
        return True

    def list_files(self, prefix):
        self.operations["list"] += 1
        return [{"object_name": name, "size": len(data)} for name, data in self.objects.items() if name.startswith(prefix)]

    def move_object(self, source, destination):
        self.operations["move"] += 1
        if self.fail_move or source not in self.objects:
            return False
        self.objects[destination] = self.objects.pop(source)
        return True

    def download_file(self, object_name, path):
        self.operations["download"] += 1
        if self.fail_download or object_name not in self.objects:
            return False
        Path(path).write_bytes(self.objects[object_name])
        return True

    def get_presigned_url(self, object_name):
        return "https://offline-storage.invalid/" + object_name


def make_docx(marker):
    document = Document()
    document.add_paragraph(marker)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx(marker):
    workbook = Workbook()
    workbook.active.title = "OfflineFixture"
    workbook.active.append(["label", "value"])
    workbook.active.append([marker, 42])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def make_pdf(marker):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
    })
    content = DecodedStreamObject()
    content.set_data(f"BT /F1 12 Tf 72 720 Td ({marker}) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(content)
    buffer = io.BytesIO()
    writer.write(buffer)
    writer.close()
    return buffer.getvalue()


class ChatUploadRegression(unittest.TestCase):
    def setUp(self):
        self.started = time.monotonic()
        self.storage = MemoryStorage()
        self.records = {}
        self.model_prompts = []
        self.network_attempts = 0
        self.details = {}

        def no_network(*args, **kwargs):
            self.network_attempts += 1
            raise AssertionError("network access is forbidden in the chat upload test")

        network_context = contextlib.ExitStack()
        network_context.enter_context(patch.object(socket, "create_connection", no_network))
        network_context.enter_context(patch.object(socket.socket, "connect", no_network))
        self.addCleanup(network_context.close)

        async def fixed_model(message, *args, **kwargs):
            self.model_prompts.append(message)
            yield {"type": "answer_replace", "content": FIXED_ANSWER}
            yield {"type": "complete", "result": {"answer": FIXED_ANSWER, "raw_answer": FIXED_ANSWER, "sources": []}}

        async def no_recommendations(*args, **kwargs):
            return []

        def load_history(conversation_id, user):
            return self.records.setdefault(conversation_id, {"messages": []}), conversation_id

        namespace = {
            "__name__": "offline_chat_route_functions",
            "Path": Path, "List": List, "Optional": Optional, "Dict": Dict, "Any": Any,
            "UploadFile": UploadFile, "File": File, "Request": Request,
            "JSONResponse": JSONResponse, "StreamingResponse": StreamingResponse,
            "io": io, "re": re, "ET": ET, "zipfile": zipfile, "PdfReader": PdfReader,
            "subprocess": subprocess, "tempfile": tempfile, "uuid": uuid, "json": json,
            "asyncio": asyncio, "sys": sys, "mimetypes": mimetypes,
            "base64": base64, "shutil": shutil,
            "MAX_FILE_TEXT_CHARS": 12000, "MAX_HISTORY_ROUNDS": 6, "FILE_HISTORY_SNIPPET_CHARS": 3000,
            "IMAGE_SUFFIXES_CHAT": {".png", ".jpg", ".jpeg", ".bmp", ".gif"},
            "storage_service": self.storage, "STAGING_PREFIX": "staging",
            "GRAPH_LOCAL_OBJECT_PREFIX": "graph-local://", "GRAPH_LOCAL_UPLOAD_DIRS": {},
            "GRAPH_LOCAL_FILE_ID_PREFIX": "graph_local",
            "now_ms": lambda: 1700000000000, "now_display": lambda: "offline-time",
            "today_str": lambda: "offline-date", "safe_segment": lambda value: str(value),
            "get_logged_in_user": lambda: {"role": "guest", "name": "offline-fixture"},
            "build_user_brief": lambda user, conversation_id: {"role": "guest", "name": "offline-fixture"},
            "load_history_record": load_history,
            "save_history_record": lambda record, key: self.records.__setitem__(key, record),
            "_active_streams": {}, "iterate_chat_events": fixed_model,
            "generate_recommendations": no_recommendations,
            "env_int": lambda name, default: 0,
        }
        tree = ast.parse(SOURCE.read_text())
        selected = []
        for item in tree.body:
            if isinstance(item, (ast.FunctionDef, ast.AsyncFunctionDef)) and item.name in FUNCTION_NAMES:
                item.decorator_list = []
                selected.append(item)
        self.assertEqual({item.name for item in selected}, FUNCTION_NAMES)
        exec(compile(ast.Module(body=selected, type_ignores=[]), str(SOURCE), "exec"), namespace)
        self.namespace = namespace
        application = FastAPI()
        application.add_api_route("/api/upload", namespace["upload_file"], methods=["POST"])
        application.add_api_route("/api/chat", namespace["chat_endpoint"], methods=["POST"])
        self.runner = asyncio.Runner()
        self.async_client = httpx.AsyncClient(
            transport=httpx.ASGITransport(app=application), base_url="http://test",
        )
        self.client = types.SimpleNamespace(
            post=lambda *args, **kwargs: self.runner.run(self.async_client.post(*args, **kwargs)),
        )
        self.addCleanup(self.runner.close)
        self.addCleanup(lambda: self.runner.run(self.async_client.aclose()))

    def tearDown(self):
        self.assertEqual(self.network_attempts, 0)
        self.assertEqual(self.namespace["_active_streams"], {})
        CASE_DETAILS.append({
            "case": self.id().rsplit(".", 1)[-1],
            "elapsed_ms": round((time.monotonic() - self.started) * 1000, 3),
            "storage_operations": dict(self.storage.operations),
            "model_stub_calls": len(self.model_prompts),
            "network_attempts": self.network_attempts,
            **self.details,
        })

    def upload(self, name, contents, mime):
        response = self.client.post("/api/upload", files={"files": (name, contents, mime)})
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["code"], 0)
        self.assertEqual(body["data"]["files"][0]["filename"], name)
        return body["data"]["files"][0]["file_id"]

    def chat_response(self, file_id):
        with contextlib.redirect_stderr(io.StringIO()):
            response = self.client.post("/api/chat", json={
                "message": "Summarize the synthetic fixture.", "conversation_id": "offline-conversation",
                "file_ids": [file_id], "web_search": False, "enable_thinking": False,
            })
        return response

    def chat_error(self, file_id, reason_fragment, message="文件解析失败"):
        response = self.chat_response(file_id)
        self.assertEqual(response.status_code, 500)
        body = response.json()
        self.assertEqual(body["code"], 1)
        self.assertEqual(body["msg"], message)
        self.assertIn(reason_fragment, body["data"]["reason"])
        self.assertEqual(self.namespace["_active_streams"], {})
        self.assertEqual(self.model_prompts, [])
        self.assertEqual(self.records["offline-conversation"]["messages"], [])
        return body

    def chat(self, file_id):
        response = self.chat_response(file_id)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.headers["content-type"].startswith("text/event-stream"))
        events = [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith("data: ")]
        self.assertNotIn("error", [event["type"] for event in events])
        self.assertEqual(events[0]["type"], "stream_open")
        self.assertEqual(events[-1]["type"], "done")
        self.assertEqual(events[-1]["data"]["answer"], FIXED_ANSWER)
        self.assertEqual(len(self.records["offline-conversation"]["messages"]), 1)
        return events[-1]["data"]

    def successful_fixture(self, suffix, contents, mime, marker):
        file_id = self.upload("fixture" + suffix, contents, mime)
        data = self.chat(file_id)
        self.assertEqual(len(data["file_contexts"]), 1)
        context = data["file_contexts"][0]
        self.assertEqual(context["kind"], "text")
        self.assertIn(marker, context["text"])
        self.assertIn(marker, self.model_prompts[0])
        self.details = {
            "upload_http_status": 200, "chat_http_status": 200, "terminal_sse_event": "done",
            "format": suffix, "real_parser_extracted_marker": True,
            "marker_in_model_prompt": True, "parsed_characters": len(context["text"]),
        }

    def test_markdown_upload_chat(self):
        marker = "CHAT_MD_FIXTURE"
        self.successful_fixture(".md", ("# Heading\n" + marker).encode(), "text/markdown", marker)

    def test_docx_upload_chat(self):
        marker = "CHAT_DOCX_FIXTURE"
        self.successful_fixture(".docx", make_docx(marker), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", marker)

    def test_xlsx_upload_chat(self):
        marker = "CHAT_XLSX_FIXTURE"
        self.successful_fixture(".xlsx", make_xlsx(marker), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", marker)

    def test_pdf_upload_chat(self):
        marker = "CHAT_PDF_FIXTURE"
        self.successful_fixture(".pdf", make_pdf(marker), "application/pdf", marker)

    def test_corrupt_pdf_returns_http_500(self):
        file_id = self.upload("corrupt.pdf", b"%PDF-1.4\nbroken fixture\n", "application/pdf")
        body = self.chat_error(file_id, "内容为空")
        self.assertIn("corrupt.pdf", body["data"]["reason"])
        self.details = {
            "upload_http_status": 200, "chat_http_status": 500,
            "explicit_parse_error_returned": True, "stream_lock_released": True,
        }

    def test_storage_upload_unavailable_returns_http_500(self):
        self.storage.fail_upload = True
        response = self.client.post("/api/upload", files={"files": ("fixture.md", b"offline", "text/markdown")})
        self.assertEqual(response.status_code, 500)
        body = response.json()
        self.assertEqual(body["code"], 1)
        self.assertIn("文件存储失败", body["data"]["reason"])
        self.assertEqual(self.model_prompts, [])
        self.assertEqual(self.storage.objects, {})
        self.details = {"upload_http_status": 500, "explicit_storage_error_returned": True}

    def test_storage_download_unavailable_returns_http_500(self):
        file_id = self.upload("fixture.md", b"CHAT_DOWNLOAD_FIXTURE", "text/markdown")
        self.storage.fail_download = True
        body = self.chat_error(file_id, "下载失败")
        self.assertIn("fixture.md", body["data"]["reason"])
        self.details = {
            "upload_http_status": 200, "chat_http_status": 500,
            "explicit_download_error_returned": True, "stream_lock_released": True,
        }

    def test_pdf_fallback_timeout_is_bounded_and_returns_http_500(self):
        file_id = self.upload("fallback.pdf", make_pdf("CHAT_FALLBACK_FIXTURE"), "application/pdf")
        calls = []

        def timed_out(command, **kwargs):
            calls.append(kwargs.get("timeout"))
            raise subprocess.TimeoutExpired(command, kwargs.get("timeout") or 0)

        with patch.dict(self.namespace, {"PdfReader": None, "subprocess": types.SimpleNamespace(run=timed_out)}):
            self.chat_error(file_id, "内容为空")
        self.assertEqual(calls, [60])
        self.details = {
            "upload_http_status": 200, "chat_http_status": 500,
            "pdftotext_timeout_seconds": calls[0], "timeout_injected_without_waiting": True,
            "explicit_parse_error_returned": True, "stream_lock_released": True,
        }

    def test_failed_parse_releases_stream_for_same_conversation_retry(self):
        file_id = self.upload("blank.md", b" \n\t ", "text/markdown")
        staging_object = self.namespace["resolve_staging_files"]([file_id])[0]["staging_object"]
        self.chat_error(file_id, "内容为空")
        self.assertIn(staging_object, self.storage.objects)
        self.assertEqual(self.storage.operations["move"], 0)
        retry_id = self.upload("retry.md", b"CHAT_RETRY_FIXTURE", "text/markdown")
        data = self.chat(retry_id)
        self.assertIn("CHAT_RETRY_FIXTURE", data["file_contexts"][0]["text"])
        self.assertEqual(len(self.model_prompts), 1)
        self.details = {
            "first_chat_http_status": 500, "retry_chat_http_status": 200,
            "same_conversation_retry_succeeded": True, "stream_lock_released": True,
            "failed_file_staging_preserved": True,
        }

    def test_text_reader_exception_returns_http_500(self):
        file_id = self.upload("unreadable.md", b"CHAT_READ_ERROR_FIXTURE", "text/markdown")
        with patch.object(Path, "read_text", side_effect=OSError("offline read failure")):
            self.chat_error(file_id, "offline read failure")
        self.details = {
            "upload_http_status": 200, "chat_http_status": 500,
            "parser_exception_propagated": True, "stream_lock_released": True,
        }

    def test_download_failure_retries_with_original_file_id(self):
        file_id = self.upload("original.md", b"CHAT_ORIGINAL_ID_FIXTURE", "text/markdown")
        staging_object = self.namespace["resolve_staging_files"]([file_id])[0]["staging_object"]
        self.storage.fail_download = True
        self.chat_error(file_id, "下载失败")
        self.assertIn(staging_object, self.storage.objects)
        self.assertEqual(self.storage.operations["move"], 0)
        self.storage.fail_download = False
        data = self.chat(file_id)
        self.assertIn("CHAT_ORIGINAL_ID_FIXTURE", data["file_contexts"][0]["text"])
        self.assertEqual(data["uploaded_files"][0]["file_id"], file_id)
        self.assertNotIn(staging_object, self.storage.objects)
        self.details = {
            "first_chat_http_status": 500, "retry_chat_http_status": 200,
            "original_file_id_retry_succeeded": True, "staging_preserved_until_success": True,
        }

    def test_graph_local_file_uses_validated_local_source(self):
        with tempfile.TemporaryDirectory(prefix="chat_graph_fixture_") as directory:
            upload_dir = Path(directory)
            local_file = upload_dir / "graph.md"
            local_file.write_text("CHAT_GRAPH_LOCAL_FIXTURE")
            encoded_name = base64.urlsafe_b64encode(local_file.name.encode()).decode().rstrip("=")
            file_id = "graph_local:internal_binding:" + encoded_name
            with patch.dict(self.namespace, {"GRAPH_LOCAL_UPLOAD_DIRS": {"internal_binding": upload_dir}}):
                data = self.chat(file_id)
            self.assertIn("CHAT_GRAPH_LOCAL_FIXTURE", data["file_contexts"][0]["text"])
            self.assertEqual(self.storage.operations["download"], 0)
            self.assertTrue(local_file.is_file())
        self.details = {
            "chat_http_status": 200, "real_graph_local_id_decoder": True,
            "validated_local_source_parsed": True, "remote_download_calls": 0,
        }

    def test_graph_local_source_rejects_path_outside_allowed_directory(self):
        with tempfile.TemporaryDirectory(prefix="chat_graph_scope_") as directory:
            root = Path(directory)
            allowed = root / "allowed"
            allowed.mkdir()
            outside = root / "outside.md"
            outside.write_text("OUTSIDE_FIXTURE_MUST_NOT_BE_PARSED")
            escape = allowed / "escape.md"
            escape.symlink_to(outside)
            with patch.dict(self.namespace, {"GRAPH_LOCAL_UPLOAD_DIRS": {"internal_binding": allowed}}):
                with self.assertRaisesRegex(RuntimeError, "文件来源无效"):
                    self.namespace["build_uploaded_file_contexts"]([{
                        "filename": "escape.md", "object_name": "graph-local://" + str(escape),
                    }])
            self.assertEqual(self.storage.operations["download"], 0)
        self.details = {"symlink_escape_rejected": True, "model_stub_calls": 0}

    def test_move_failure_preserves_staging_and_releases_stream(self):
        file_id = self.upload("move-failure.md", b"CHAT_MOVE_FAILURE_FIXTURE", "text/markdown")
        staging_object = self.namespace["resolve_staging_files"]([file_id])[0]["staging_object"]
        self.storage.fail_move = True
        self.chat_error(file_id, "文件保存失败", message="文件保存失败")
        self.assertIn(staging_object, self.storage.objects)
        self.assertEqual(self.storage.operations["download"], 1)
        self.assertEqual(self.storage.operations["move"], 1)
        self.details = {
            "chat_http_status": 500, "staging_preserved": True,
            "stream_lock_released": True, "model_stub_calls": 0,
        }


if __name__ == "__main__":
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(ChatUploadRegression)
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    summary = {
        "source_root": str(ROOT), "source_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        "entrypoints": ["POST /api/upload", "POST /api/chat"],
        "original_function_bodies": sorted(FUNCTION_NAMES),
        "cases": CASE_DETAILS,
        "total": result.testsRun, "failures": len(result.failures), "errors": len(result.errors),
        "passed": result.wasSuccessful(),
        "stubbed_boundaries": [
            "Object storage: in-memory byte dictionary; downloads write synthetic bytes into temporary parser directories.",
            "Login identity: fixed synthetic guest; no real session, permission, Redis, or MySQL calls.",
            "Conversation history: in-memory dictionary.",
            "LLM/agent events and recommendations: fixed deterministic outputs; model prompts are captured for assertions.",
            "Graph-local directory allowlist: temporary fixture directories; the original ID decoder and path checks run unchanged.",
            "PDF fallback timeout case only: injected TimeoutExpired; normal PDF uses real pypdf parsing.",
        ],
        "safety": {"external_network_requests": 0, "real_model_requests": 0, "real_database_operations": 0, "production_app_imported": False},
        "limitations": [
            "This tests original backend route functions through in-process httpx ASGITransport in an isolated FastAPI app, not a deployed service or its authentication middleware.",
            "SSE completion with fixed model output does not measure actual model latency or answer quality.",
            "Corrupt PDF, failed download, empty text, parser exceptions, and PDF fallback timeouts must return explicit errors without invoking the model.",
        ],
    }
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(summary, ensure_ascii=False, indent=2))
    REPORT.chmod(0o600)
    print(json.dumps({"report": str(REPORT), "tests": result.testsRun, "passed": result.wasSuccessful()}))
    raise SystemExit(0 if result.wasSuccessful() else 1)
